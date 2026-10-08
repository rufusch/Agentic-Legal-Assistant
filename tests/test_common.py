import hashlib
import time
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient

from backend.app import create_app


@pytest.fixture
def client(tmp_path):
    app = create_app(tmp_path, {'alice': 'tenant-a', 'bob': 'tenant-b'}, review_engine='extractive')
    with TestClient(app) as client:
        client.headers['Authorization'] = 'Bearer alice'
        yield client
    app.state.db.close()


def upload(client, text=b'Payment is due on 15 October.\nTermination requires notice.', key='upload-1'):
    body = {'file_name': 'contract.txt', 'content_type': 'text/plain', 'size_bytes': len(text), 'sha256': hashlib.sha256(text).hexdigest()}
    response = client.post('/api/v1/documents/uploads', json=body, headers={'Idempotency-Key': key})
    assert response.status_code == 201, response.text
    data = response.json()['data']
    assert client.put(data['upload_url'], content=text).status_code == 200
    return data, body


def finish(client, data):
    response = client.post(f"/api/v1/documents/{data['document_id']}/complete", json={'upload_id': data['upload_id'], 'metadata': {'jurisdiction': 'IN'}}, headers={'Idempotency-Key': 'complete-1'})
    assert response.status_code == 202, response.text
    result = response.json()['data']
    if client.app.state.worker.thread:
        deadline = time.monotonic() + 40
        while time.monotonic() < deadline:
            state = client.get(f"/api/v1/jobs/{result['job_id']}").json()['data']
            if state['status'] in {'completed', 'completed_with_warnings', 'failed', 'cancelled'}:
                break
            time.sleep(.1)
        else:
            raise AssertionError('Processing timed out in integration test')
    return result


def test_full_lifecycle_and_exact_source(client):
    data, _ = upload(client)
    result = finish(client, data)
    doc = client.get(f"/api/v1/documents/{data['document_id']}").json()['data']
    assert doc['status'] == 'ready'
    assert doc['page_count'] == 1
    listing = client.get('/api/v1/documents?status=ready&query=contract').json()['data']
    assert listing['items'][0]['id'] == doc['id']
    db = client.app.state.db
    chunk_id = db.execute("SELECT id FROM resources WHERE kind='chunk'").fetchone()[0]
    chunk = client.get(f"/api/v1/documents/{doc['id']}/chunks/{chunk_id}").json()['data']
    preview = client.get(chunk['preview_url'])
    assert preview.content.decode()[chunk['start_offset']:chunk['end_offset']] == chunk['text']
    assert b'Payment is due' not in (db.execute("SELECT body FROM resources WHERE kind='chunk'").fetchone()[0])
    assert client.get(f"/api/v1/jobs/{result['job_id']}").json()['data']['status'] == 'completed'
    events = client.get(result['events_url']).text
    assert 'job.snapshot' in events and 'job.completed' in events
    resumed = client.get(result['events_url'], headers={'Last-Event-ID': '2'}).text
    assert 'id: 1\n' not in resumed and 'id: 2\n' not in resumed
    assert 'job.completed' in resumed
    assert client.post(f"/api/v1/jobs/{result['job_id']}/cancel").json()['data']['status'] == 'completed'


def test_auth_and_tenant_isolation(client):
    data, _ = upload(client)
    result = finish(client, data)
    assert client.get('/api/v1/home', headers={'Authorization': ''}).status_code == 401
    for url in [f"/api/v1/documents/{data['document_id']}", f"/api/v1/jobs/{result['job_id']}", result['events_url']]:
        assert client.get(url, headers={'Authorization': 'Bearer bob'}).status_code == 404
    assert client.get('/api/v1/documents', headers={'Authorization': 'Bearer bob'}).json()['data']['items'] == []


def test_idempotency(client):
    data, body = upload(client)
    replay = client.post('/api/v1/documents/uploads', json=body, headers={'Idempotency-Key': 'upload-1'})
    assert replay.json()['data'] == data
    body['file_name'] = 'different.txt'
    assert client.post('/api/v1/documents/uploads', json=body, headers={'Idempotency-Key': 'upload-1'}).status_code == 409
    first = finish(client, data)
    second = client.post(f"/api/v1/documents/{data['document_id']}/complete", json={'upload_id': data['upload_id'], 'metadata': {'jurisdiction': 'IN'}}, headers={'Idempotency-Key': 'complete-1'})
    assert second.json()['data'] == first


def test_integrity_validation_and_state(client):
    data, _ = upload(client)
    assert client.put(data['upload_url'], content=b'x').status_code == 422
    assert client.put(data['upload_url'], content=b'x' * 100).status_code == 413
    assert client.put(data['upload_url'].split('?')[0] + '?token=bad', content=b'x').status_code == 401
    assert client.post(f"/api/v1/documents/{data['document_id']}/complete", json={'upload_id': str(uuid4())}).status_code == 404
    finish(client, data)
    assert client.put(data['upload_url'], content=b'Payment is due on 15 October.\nTermination requires notice.').status_code == 409
    assert client.get('/api/v1/documents?limit=0').status_code == 400
    assert client.get('/api/v1/documents/not-uuid').status_code == 400


def test_bad_content_fails_visibly(client):
    data, _ = upload(client, b'\xff\xfeinvalid')
    result = finish(client, data)
    doc = client.get(f"/api/v1/documents/{data['document_id']}").json()['data']
    assert doc['status'] == 'failed' and doc['failure']['code'] == 'INVALID_ENCODING'
    assert 'job.failed' in client.get(result['events_url']).text


def test_pagination_and_invalid_sse_cursor(client):
    first, _ = upload(client)
    upload(client, b'Second document', 'upload-2')
    page = client.get('/api/v1/documents?limit=1').json()['data']
    assert page['next_cursor']
    second = client.get('/api/v1/documents?limit=1&cursor=' + page['next_cursor']).json()['data']
    assert second['items'][0]['id'] != page['items'][0]['id']
    result = finish(client, first)
    assert client.get(result['events_url'], headers={'Last-Event-ID': 'bad'}).status_code == 400


def test_cancel_pending_processing(tmp_path):
    app = create_app(tmp_path, {'alice': 'tenant-a'}, start_worker=False)
    with TestClient(app, headers={'Authorization': 'Bearer alice'}) as client:
        data, _ = upload(client)
        result = finish(client, data)
        cancelled = client.post(f"/api/v1/jobs/{result['job_id']}/cancel").json()['data']
        assert cancelled['status'] == 'cancelled'
        assert client.get(f"/api/v1/documents/{data['document_id']}").json()['data']['failure']['code'] == 'CANCELLED'
        assert 'cancelled' in client.get(result['events_url']).text
    app.state.store.close()


def test_restart_resumes_interrupted_job_and_preserves_data(tmp_path):
    app = create_app(tmp_path, {'alice': 'tenant-a'}, start_worker=False)
    with TestClient(app, headers={'Authorization': 'Bearer alice'}) as client:
        data, _ = upload(client)
        result = finish(client, data)
    app.state.db.close()
    recovered = create_app(tmp_path, {'alice': 'tenant-a'})
    with TestClient(recovered, headers={'Authorization': 'Bearer alice'}) as client:
        job = client.get(f"/api/v1/jobs/{result['job_id']}").json()['data']
        for _ in range(200):
            job = client.get(f"/api/v1/jobs/{result['job_id']}").json()['data']
            if job['status'] == 'completed':
                break
            time.sleep(.1)
        assert job['status'] == 'completed'
        assert client.get(f"/api/v1/documents/{data['document_id']}").json()['data']['status'] == 'ready'
    recovered.state.db.close()


def test_docx_parsing_includes_tables(client):
    import io
    from docx import Document
    source = Document()
    source.add_paragraph('Agreement terms')
    source.add_table(rows=1, cols=1).cell(0, 0).text = 'Payment on 15 October'
    buffer = io.BytesIO()
    source.save(buffer)
    raw = buffer.getvalue()
    body = {'file_name': 'agreement.docx', 'content_type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'size_bytes': len(raw), 'sha256': hashlib.sha256(raw).hexdigest()}
    data = client.post('/api/v1/documents/uploads', json=body).json()['data']
    assert client.put(data['upload_url'], content=raw).status_code == 200
    finish(client, data)
    assert client.get(f"/api/v1/documents/{data['document_id']}").json()['data']['status'] == 'ready'
    chunk_id = client.app.state.db.execute("SELECT id FROM resources WHERE kind='chunk'").fetchone()[0]
    text = client.get(f"/api/v1/documents/{data['document_id']}/chunks/{chunk_id}").json()['data']['text']
    assert 'Agreement terms' in text and 'Payment on 15 October' in text

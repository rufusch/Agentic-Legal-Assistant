from backend.official_sources import citation_source_text
from backend.context_budget import source_budget
"""Legal Drafting API, durable workflow jobs and version/consent boundaries."""
import hashlib
import json
import os
import time
from dataclasses import replace
from datetime import datetime, timezone
from io import BytesIO
from typing import Literal
from uuid import UUID, uuid4

from fastapi import Request
from fastapi.responses import Response

from backend.drafting_schema import CreateDraft, Answers, GenerateDraft, EditSection, DraftFeedback
from backend.drafting_requirements import requirements, missing, refresh_status, prefill
from backend.drafting_engine import generate
from backend.jobs import TERMINAL
from backend.models.review_llm import LocalReviewLLM, ReviewModelError, ReviewCancelled
from backend.parsing import chunks as split_chunks
from backend.retrieval import search
from backend.review_engine import confidence, warning


def uid(): return str(uuid4())
def now(): return datetime.now(timezone.utc).isoformat()


def install(app, store, worker, get, envelope, APIError, idempotent, drafting_llm=None, verify_llm=None):
    # A separate instance/configuration is deliberate: training/replacing drafting never changes review.
    def env(name, fallback): return os.getenv('LEXIMIND_DRAFTING_'+name, fallback)
    llm = drafting_llm or LocalReviewLLM.for_role('drafting')
    from backend.agent_loop import verifier, verification_info
    checker = verifier(llm, verify_llm)
    app.state.drafting_llm = llm; app.state.drafting_verify_llm = checker
    def metadata(): return {**llm.metadata,'id':'leximind-drafting','workflow':'drafting','output_schema':'legal-draft-v1','prompt_version':'grounded-drafting-v1','verifier':verification_info(llm,checker)}

    def snapshot(tenant, draft):
        store.save(tenant,'draft_version',{'id':uid(),'draft_id':draft['id'],'version':draft['version'],'source_document_ids':draft['source_document_ids'],'created_at':now(),'draft':draft})

    def active(draft):
        return draft['status'] in {'checking_requirements','retrieving','drafting','verifying','queued'}

    def ensure_idle(draft):
        if active(draft): raise APIError(409,'DRAFT_RUNNING','Wait for or cancel the active drafting job.')

    def queue(tenant, draft, phase):
        if sum(j['status'] not in TERMINAL for j in store.all(tenant,'job')) >= 20: raise APIError(429,'QUEUE_FULL','Wait for queued jobs.',True)
        jid=uid();draft['job_id']=jid;draft['status']='checking_requirements' if phase=='requirements' else 'queued';draft['failure']=None
        job={'id':jid,'workflow':'drafting','resource_id':draft['id'],'phase':phase,'status':draft['status'],'progress':0,'created_at':now(),'result_id':None,'warnings':[],'failure':None}
        store.save(tenant,'draft',draft);store.save(tenant,'job',job)
        store.db.execute('INSERT INTO queue VALUES(?,?,?,?)',(jid,tenant,draft['id'],'queued'))
        store.event(jid,'job.progress',status=job['status'],progress=0,message='Legal Drafting '+phase+' queued')
        worker.wake.set()
        return {'draft_id':draft['id'],'job_id':jid,'status':draft['status'],'events_url':f'/api/v1/jobs/{jid}/events'}

    def capture_intake(tenant, draft):
        text='USER-PROVIDED INTAKE — unverified assertions and drafting instructions, not independent findings.\n'
        fields={'document_type':draft['document_type'],'jurisdiction':draft['jurisdiction'],'court':draft['court'],'instructions':draft['instructions'],**{'facts.'+str(k):v for k,v in draft['facts'].items()}}
        text+='\n'.join(f'{key}: {json.dumps(value,ensure_ascii=False)}' for key,value in fields.items() if value is not None)
        rid=uid();raw=text.encode('utf-8');created=now()
        existing=store.all(tenant,'document')
        if len(existing)>=200 or sum(d['size_bytes'] for d in existing)+len(raw)>500*1024*1024:raise APIError(429,'STORAGE_QUOTA','Workspace storage limit reached.')
        doc={'id':rid,'name':f'Legal Drafting intake v{draft["version"]}.txt','media_type':'text/plain','size_bytes':len(raw),'sha256':hashlib.sha256(raw).hexdigest(),'status':'ready','created_at':created,'metadata':{'document_type':'user_input','jurisdiction':draft['jurisdiction'],'draft_id':draft['id']},'warnings':[],'ocr_used':False,'page_count':None,'parser_version':'user-intake-v1','index_version':'bm25-lsa-rrf-v1','processing':{'stage':'ready','progress':1},'failure':None,'uploaded':True,'retention_expires_at':time.time()+worker.retention_days*86400}
        parts=[]
        for start,end,value in split_chunks(text):
            parts.append({'id':uid(),'document_id':rid,'text':value,'start_offset':start,'end_offset':end,'page':None,'source_part':1,'section':'User-provided intake','ocr':False})
        doc['chunk_count']=len(parts)
        store.write_blob(rid,raw);store.save(tenant,'document',doc)
        for chunk in parts:store.save(tenant,'chunk',chunk)
        draft['context_document_id']=rid
        return doc,parts

    def retrieve(tenant,draft):
        intake=store.get(tenant,'document',draft.get('context_document_id'))
        if not intake: intake,_=capture_intake(tenant,draft)
        all_docs=store.all(tenant,'document');docs={d['id']:d for d in all_docs if d['status']=='ready'}
        selected=set(draft['supporting_document_ids']);selected.add(intake['id'])
        if any(rid not in docs for rid in selected):raise ReviewModelError('SOURCES_NOT_READY','A selected source is unavailable; update the matter sources.',False)
        authority_ids={d['id'] for d in docs.values() if d['metadata'].get('document_type') in {'statute','judgment'} and d['metadata'].get('jurisdiction') in {draft['jurisdiction'],draft['jurisdiction'].split('-')[0]}}
        style_ids={d['id'] for d in docs.values() if d['metadata'].get('document_type')=='past_draft'}
        case_ids=selected-style_ids
        all_chunks=store.all(tenant,'chunk')
        query=draft['document_type'].replace('_',' ')+' '+draft['instructions']+' '+json.dumps(draft['facts'],ensure_ascii=False)
        case_pool=[c for c in all_chunks if c['document_id'] in case_ids and c['document_id']!=intake['id']]
        authority_pool=[c for c in all_chunks if c['document_id'] in authority_ids or c['document_id'] in selected and docs[c['document_id']]['metadata'].get('document_type') in {'statute','judgment'}]
        # Selected authorities with conflicting explicit jurisdiction are never silently treated as governing law.
        authority_pool=[c for c in authority_pool if docs[c['document_id']]['metadata'].get('jurisdiction') in {draft['jurisdiction'],draft['jurisdiction'].split('-')[0]}]
        if app.state.official_sources.enabled:
            current_ids={d['id'] for d in app.state.official_sources.current_documents(tenant)}
            authority_pool=[c for c in authority_pool if c['document_id'] in current_ids]
        authority_candidates={c['id'] for c in authority_pool}
        case_pool=[c for c in case_pool if docs[c['document_id']]['metadata'].get('document_type') not in {'statute','judgment'}]
        chosen=[c for c in all_chunks if c['document_id']==intake['id']]
        ranked=[*search(case_pool,query,12),*search(authority_pool,query,10)]
        used=sum(len(c['text']) for c in chosen)
        budget=source_budget(llm,24000)
        if used>budget:raise ReviewModelError('DRAFT_CONTEXT_LIMIT',f'Intake exceeds the {budget:,}-character drafting context budget. Narrow the supplied intake.',False)
        for c in ranked:
            if used+len(c['text'])<=budget and c['id'] not in {v['id'] for v in chosen}:chosen.append(c);used+=len(c['text'])
        style_hits=search([c for c in all_chunks if c['document_id'] in style_ids],query,3)
        examples=[{'document_name':docs[c['document_id']]['name'],'style_only':True,'text':c['text'][:1200]} for c in style_hits]
        source_ids={c['document_id'] for c in chosen}
        draft['source_document_ids']=sorted(source_ids | selected | {c['document_id'] for c in style_hits})
        draft['retrieval']={'method':'bm25-lsa-rrf-v1','case_chunks_available':len(case_pool),'authority_chunks_available':len(authority_pool),'selected_evidence_chunks':len(chosen),'selected_authority_chunks':sum(c['id'] in authority_candidates for c in chosen),'style_examples':len(examples),'context_characters':used,'scope':'Relevant excerpts from selected case sources and jurisdiction-filtered tenant authorities. Style examples are not factual or legal evidence. No external legal database.'}
        return [docs[rid] for rid in source_ids],chosen,examples

    def execute(tenant, rid, jid):
        draft=get(tenant,'draft',rid);job=get(tenant,'job',jid)
        try:
            def progress(stage,value,message):
                with store.transaction():
                    if worker.cancelled(tenant,jid):raise ReviewCancelled()
                    draft['status']=stage;job.update(status=stage,progress=value)
                    store.save(tenant,'draft',draft);store.save(tenant,'job',job)
                    store.event(jid,'job.progress',status=stage,progress=value,message=message)
            if job['phase']=='requirements':
                progress('checking_requirements',.5,'Checking intake fields and recording unresolved information')
                sources=[get(tenant,'document',i) for i in draft['supporting_document_ids']]
                draft['requirements']=requirements(draft);prefill(draft,sources,[c for c in store.all(tenant,'chunk') if c['document_id'] in set(draft['supporting_document_ids'])]);refresh_status(draft)
                draft['warnings']=[warning('missing_information','Intake checklist scope','This is a drafting intake checklist, not a certified jurisdictional filing checklist. Supporting evidence is strongly recommended.','info')]
            else:
                progress('retrieving',.15,'Retrieving relevant case evidence, tenant authorities and style examples')
                official=app.state.official_sources.enrich(tenant,draft['instructions']+' '+draft['document_type'],lambda:worker.cancelled(tenant,jid))
                with store.transaction():
                    documents,chunks,examples=retrieve(tenant,draft)
                    draft['retrieval']['official_sources']=official
                    store.save(tenant,'draft',draft)
                result,training=generate(draft,documents,chunks,examples,llm,lambda:worker.cancelled(tenant,jid),progress,verify_edits=job['phase']=='verify',verify_llm=checker)
                result['model']={**result['model'],**metadata()};draft=result
                # Only the sanitized published result can become an approved training target.
                from backend.drafting_training import target_from_draft
                training['target']=target_from_draft(draft,chunks)
            with store.transaction():
                if worker.cancelled(tenant,jid):return
                store.save(tenant,'draft',draft);snapshot(tenant,draft)
                if job['phase']!='requirements':store.save(tenant,'draft_training',{'id':uid(),'draft_id':rid,'version':draft['version'],'source_document_ids':draft['source_document_ids'],'messages':training['messages'],'target':training['target']})
                job.update(status='completed_with_warnings',progress=1,result_id=rid,warnings=draft['warnings'])
                store.save(tenant,'job',job);store.event(jid,'job.completed',status=job['status'],result_id=rid,resource_status=draft['status'])
                store.db.execute('DELETE FROM queue WHERE job=?',(jid,));store.audit(tenant,'draft.'+job['phase'],rid,model=draft['model'])
        except ReviewCancelled:return
        except Exception as exc:
            with store.transaction():
                if worker.cancelled(tenant,jid):return
                failure={'code':exc.code,'message':exc.message,'retryable':exc.retryable} if isinstance(exc,ReviewModelError) else {'code':'DRAFT_VERIFICATION_FAILED','message':'Draft could not pass source verification. No completed draft was published.','retryable':True}
                draft.update(status='failed',failure=failure);job.update(status='failed',failure=failure)
                store.save(tenant,'draft',draft);store.save(tenant,'job',job);store.event(jid,'job.failed',**failure);store.db.execute('DELETE FROM queue WHERE job=?',(jid,))

    worker.handlers['drafting']=execute

    @app.get('/api/v1/drafting/config')
    def config(request: Request):return envelope(request,{**llm.status(),'model':metadata(),'document_types':['bail_application','petition','notice','affidavit','contract_clause','other'],'languages':['en'],'authority_scope':'Tenant document library; no external legal database.'})

    @app.post('/api/v1/drafts',status_code=202)
    def create(request: Request, body: CreateDraft):
        def build():
            tenant=request.state.tenant
            with store.transaction():
                for rid in body.supporting_document_ids:
                    source=get(tenant,'document',rid)
                    if source['status']!='ready':raise APIError(409,'SOURCES_NOT_READY','Supporting documents must be ready.')
                draft={**body.model_dump(mode='json'),'id':uid(),'status':'checking_requirements','title':body.document_type.replace('_',' ').title(),'created_at':now(),'version':1,'requirements':[],'sections':[],'unresolved_placeholders':[],'authorities':[],'claims':[],'citations':[],'warnings':[],'source_document_ids':list(map(str,body.supporting_document_ids)),'confidence':confidence(0,'Draft not generated.'),'model':metadata(),'failure':None,'needs_verification':False,'verification':{},'retrieval':{},'context_document_id':None,'acknowledged_requirement_ids':[]}
                return envelope(request,queue(tenant,draft,'requirements'))
        return idempotent(request,body.model_dump(mode='json'),build)

    @app.get('/api/v1/drafts')
    def listing(request: Request):return envelope(request,{'items':sorted(store.all(request.state.tenant,'draft'),key=lambda d:d['created_at'],reverse=True)[:100]})

    @app.get('/api/v1/drafts/{rid}')
    def read(rid: UUID, request: Request):return envelope(request,get(request.state.tenant,'draft',rid))

    @app.get('/api/v1/drafts/{rid}/requirements')
    def intake(rid: UUID, request: Request):
        draft=get(request.state.tenant,'draft',rid)
        return envelope(request,{'draft_id':str(rid),'status':draft['status'],'items':draft['requirements'],'missing_requirement_ids':[r['id'] for r in missing(draft['requirements'])]})

    @app.patch('/api/v1/drafts/{rid}/requirements')
    def answer(rid: UUID, request: Request, body: Answers):
        tenant=request.state.tenant
        with store.transaction():
            draft=get(tenant,'draft',rid);ensure_idle(draft)
            by_id={r['id']:r for r in draft['requirements']}
            if len({str(a.requirement_id) for a in body.answers})!=len(body.answers):raise APIError(400,'VALIDATION_ERROR','Duplicate requirement answer.')
            for answer in body.answers:
                if str(answer.requirement_id) not in by_id:raise APIError(400,'VALIDATION_ERROR','Unknown requirement id.')
                item=by_id[str(answer.requirement_id)]
                if item.get('source')=='document' and (answer.value.strip() or None)!=item['current_answer']:item.update(source='user',citation=None)
                item['current_answer']=answer.value.strip() or None
                if item['key'] in {'court','instructions','jurisdiction'}:draft[item['key']]=item['current_answer']
                else:draft['facts'][item['key']]=item['current_answer']
            if len(json.dumps(draft['facts'],ensure_ascii=False))>24000:raise APIError(413,'DRAFT_CONTEXT_LIMIT','Facts exceed 24,000 characters.')
            draft.update(version=draft['version']+1,context_document_id=None,needs_verification=bool(draft['sections']))
            refresh_status(draft);store.save(tenant,'draft',draft);snapshot(tenant,draft)
            return envelope(request,{'draft_id':str(rid),'status':draft['status'],'version':draft['version'],'items':draft['requirements']})

    @app.post('/api/v1/drafts/{rid}/generate',status_code=202)
    def start(rid: UUID, request: Request, body: GenerateDraft):
        def build():
            tenant=request.state.tenant
            with store.transaction():
                draft=get(tenant,'draft',rid);ensure_idle(draft);gaps=missing(draft['requirements'])
                if not draft['requirements']:raise APIError(409,'REQUIREMENTS_NOT_READY','Complete the missing-information check first.')
                if any(r['required'] for r in gaps):raise APIError(409,'BLOCKING_INFORMATION_REQUIRED','Supply the jurisdiction and substantive drafting instructions.')
                ids=set(map(str,body.acknowledged_requirement_ids));expected={r['id'] for r in gaps}
                if ids-expected:raise APIError(400,'VALIDATION_ERROR','Acknowledgement contains an unknown or already answered requirement.')
                if gaps and (not body.proceed_with_missing_information or ids!=expected):raise APIError(409,'MISSING_INFORMATION_ACKNOWLEDGEMENT_REQUIRED','Answer the gaps or explicitly acknowledge every unresolved non-blocking field.')
                draft.update(version=draft['version']+1,acknowledged_requirement_ids=list(ids),context_document_id=None)
                capture_intake(tenant,draft)
                draft.update(sections=[],claims=[],citations=[],authorities=[],unresolved_placeholders=[],warnings=[],needs_verification=False,verification={})
                return envelope(request,queue(tenant,draft,'generate'))
        return idempotent(request,body.model_dump(mode='json'),build)

    @app.patch('/api/v1/drafts/{rid}/sections/{sid}')
    def edit(rid: UUID, sid: UUID, request: Request, body: EditSection):
        tenant=request.state.tenant
        with store.transaction():
            draft=get(tenant,'draft',rid);ensure_idle(draft)
            if body.base_version!=draft['version']:raise APIError(409,'VERSION_CONFLICT','Reload the draft before editing.')
            section=next((s for s in draft['sections'] if s['id']==str(sid)),None)
            if not section:raise APIError(404,'NOT_FOUND','Section unavailable.')
            old_claims={cid for b in section['blocks'] for cid in b['claim_ids']}
            section['blocks']=[{'id':uid(),'kind':'paragraph','text':body.text,'claim_ids':[],'citation_ids':[],'editable':True,'statement_type':'user_edit','verification_status':'unverified'}]
            draft['claims']=[c for c in draft['claims'] if c['id'] not in old_claims]
            used={cid for s in draft['sections'] for b in s['blocks'] for cid in b['citation_ids']}
            draft['citations']=[c for c in draft['citations'] if c['id'] in used]
            draft['authorities']=[{**a,'citation_ids':[c for c in a['citation_ids'] if c in used]} for a in draft['authorities'] if any(c in used for c in a['citation_ids'])]
            draft.update(version=draft['version']+1,needs_verification=True,status='completed_with_warnings',confidence=confidence(0,'User edits have not been verified.'))
            draft['warnings'].append(warning('unsupported_claim','Unverified user edits','Edited content has no inherited claims or citations. Verify it before exporting.'))
            store.save(tenant,'draft',draft);snapshot(tenant,draft);return envelope(request,draft)

    @app.post('/api/v1/drafts/{rid}/verify',status_code=202)
    def verify(rid: UUID, request: Request):
        def build():
            with store.transaction():
                draft=get(request.state.tenant,'draft',rid);ensure_idle(draft)
                if any(r['required'] for r in missing(draft['requirements'])):raise APIError(409,'BLOCKING_INFORMATION_REQUIRED','Restore jurisdiction and substantive drafting instructions before verification.')
                if not draft['sections']:raise APIError(409,'DRAFT_NOT_READY','Generate a draft before verifying edits.')
                draft['version']+=1
                return envelope(request,queue(request.state.tenant,draft,'verify'))
        return idempotent(request,{},build)

    @app.get('/api/v1/drafts/{rid}/versions')
    def versions(rid: UUID, request: Request):
        get(request.state.tenant,'draft',rid)
        return envelope(request,{'items':sorted([v for v in store.all(request.state.tenant,'draft_version') if v['draft_id']==str(rid)],key=lambda v:v['version'])})

    @app.get('/api/v1/drafts/{rid}/export')
    def export(rid: UUID, request: Request, format: Literal['txt','docx','pdf']='docx'):
        draft=get(request.state.tenant,'draft',rid)
        if draft['status'] not in {'completed','completed_with_warnings'} or draft['needs_verification']:raise APIError(409,'DRAFT_UNVERIFIED','Finish generation and verify edited content before exporting.')
        lines=['WORKING DRAFT — HUMAN REVIEW REQUIRED',draft['title'],*['WARNING: '+w['message'] for w in draft['warnings']]]
        for section in draft['sections']:
            lines.append(section['heading'])
            for block in section['blocks']:lines.append(block['text']+' '+''.join('['+next(c['label'] for c in draft['citations'] if c['id']==cid)+']' for cid in block['citation_ids']))
        lines.append('Source references')
        lines.extend(f'[{c["label"]}] {c["document_name"]}, page {c.get("page") or "body"}, offsets {c["start_offset"]}:{c["end_offset"]}: {c["quoted_text"]}' +citation_source_text(c) for c in draft['citations'])
        stream=BytesIO()
        if format=='txt':content,media='\n\n'.join(lines).encode('utf-8'),'text/plain; charset=utf-8'
        elif format=='docx':
            from docx import Document
            doc=Document()
            for line in lines:doc.add_paragraph(line)
            doc.save(stream);content,media=stream.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        else:
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            from xml.sax.saxutils import escape
            styles=getSampleStyleSheet();story=[]
            for line in lines:story.extend([Paragraph(escape(line),styles['BodyText']),Spacer(1,8)])
            SimpleDocTemplate(stream).build(story);content,media=stream.getvalue(),'application/pdf'
        return Response(content,media_type=media,headers={'Content-Disposition':f'attachment; filename="legal-draft-v{draft["version"]}.{format}"'})

    @app.post('/api/v1/drafts/{rid}/feedback')
    def feedback(rid: UUID, request: Request, body: DraftFeedback):
        draft=get(request.state.tenant,'draft',rid)
        if body.base_version!=draft['version']:raise APIError(409,'VERSION_CONFLICT','Approve only the current draft version.')
        if draft['status'] not in {'completed','completed_with_warnings'} or draft['needs_verification']:raise APIError(409,'DRAFT_UNVERIFIED','Only a checked completed draft is eligible for approval.')
        item={'id':uid(),'draft_id':str(rid),'version':draft['version'],'source_document_ids':draft['source_document_ids'],'created_at':now(),**body.model_dump()}
        store.save(request.state.tenant,'draft_feedback',item)
        return envelope(request,{'feedback_id':item['id'],'training_eligible':body.accepted and body.use_for_training})
